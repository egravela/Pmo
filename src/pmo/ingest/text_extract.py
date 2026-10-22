"""Estrazione di testo da allegati e file (PDF, DOCX, testo)."""

from __future__ import annotations

import io
import logging

log = logging.getLogger(__name__)


def estrai_testo(data: bytes, mime_type: str, filename: str = "", max_chars: int = 4000) -> str:
    mime = (mime_type or "").lower()
    nome = (filename or "").lower()
    try:
        if mime == "application/pdf" or nome.endswith(".pdf"):
            return _pdf(data, max_chars)
        if mime.endswith("wordprocessingml.document") or nome.endswith(".docx"):
            return _docx(data, max_chars)
        if mime.startswith("text/") or nome.endswith((".txt", ".md", ".csv")):
            return data.decode("utf-8", errors="replace")[:max_chars]
    except Exception as exc:  # noqa: BLE001
        log.warning("Estrazione testo fallita per %s (%s): %s", filename, mime, exc)
    return ""


def _pdf(data: bytes, max_chars: int) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parti: list[str] = []
    totale = 0
    for page in reader.pages[:15]:
        testo = page.extract_text() or ""
        parti.append(testo)
        totale += len(testo)
        if totale >= max_chars:
            break
    return "\n".join(parti).strip()[:max_chars]


def _docx(data: bytes, max_chars: int) -> str:
    import docx

    documento = docx.Document(io.BytesIO(data))
    parti = [p.text for p in documento.paragraphs if p.text.strip()]
    for tabella in documento.tables[:10]:
        for riga in tabella.rows:
            parti.append(" | ".join(c.text.strip() for c in riga.cells))
    return "\n".join(parti).strip()[:max_chars]
